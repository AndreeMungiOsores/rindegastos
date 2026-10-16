import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Configuración de márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Estilos de fuente general
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # 1. TÍTULO PRINCIPAL
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_sub = title_p.add_run("GRUPO BLISSCORP · DIRECCIÓN DE TECNOLOGÍA\n")
    run_sub.font.size = Pt(9)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

    run_title = title_p.add_run("Informe de Envío de Casos de Prueba al ERP Sea Fácil / Niuxpro")
    run_title.font.size = Pt(18)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    desc_p = doc.add_paragraph()
    desc_p.paragraph_format.space_after = Pt(14)
    desc_p.add_run("Resumen ejecutivo y técnico del despacho de comprobantes reales vía API para validación del Registro de Compras, Anexos y Asientos Contables.")

    # 2. CUADRO DE METADATOS
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Empresa Destino:", "BLISSCORP S.A.C. (Código de Empresa: 26 — Modo de Prueba)"),
        ("Destinatarios:", "Ing. Fernando Acosta (Niuxpro) / Equipo de Contabilidad"),
        ("Remitente / Fecha:", "Andree Mungi (TI BlissCorp) — 06 de Octubre de 2026"),
        ("Resultado General:", "11 Comprobantes Despachados con Éxito (100% Recibidos con HTTP 201)")
    ]

    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.0)
        cell_val.width = Inches(4.8)

        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, top=70, bottom=70, left=100, right=100)
        set_cell_margins(cell_val, top=70, bottom=70, left=100, right=100)

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_after = Pt(0)
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(9.5)
        r_lbl.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

        p_val = cell_val.paragraphs[0]
        p_val.paragraph_format.space_after = Pt(0)
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(9.5)
        if idx == 3:
            r_val.font.bold = True
            r_val.font.color.rgb = RGBColor(0x05, 0x96, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 3. IMAGEN DE LA BANDEJA DEL ERP
    img_path = r"C:\Users\Andree\.gemini\antigravity\brain\c8dc0e33-7ed9-4277-a0ca-f094345c3f50\.user_uploaded\media_1791316005253.png"
    if os.path.exists(img_path):
        h_img = doc.add_paragraph()
        h_img.paragraph_format.space_before = Pt(8)
        h_img.paragraph_format.space_after = Pt(4)
        r_himg = h_img.add_run("1. Evidencia de Recepción en Bandeja del ERP")
        r_himg.font.bold = True
        r_himg.font.size = Pt(13)
        r_himg.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.8))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(14)
        r_cap = p_cap.add_run("Figura 1: Vista de recepción del módulo Compras v1 en Niuxpro con los 11 comprobantes registrados.")
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # 4. TABLA DE CASOS ENVIADOS
    h_cases = doc.add_paragraph()
    h_cases.paragraph_format.space_before = Pt(8)
    h_cases.paragraph_format.space_after = Pt(4)
    r_hcases = h_cases.add_run("2. Matriz de Comprobantes Despachados")
    r_hcases.font.bold = True
    r_hcases.font.size = Pt(13)
    r_hcases.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    table = doc.add_table(rows=1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Caso", "Comprobante", "Proveedor / RUC", "Importe", "Tipo / Esc.", "ID ERP", "Observación / Tratamiento"]
    col_widths = [Inches(0.6), Inches(1.1), Inches(1.5), Inches(0.9), Inches(0.8), Inches(0.6), Inches(1.5)]

    hdr_row = table.rows[0]
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, "0284C7")
        set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    cases_rows = [
        ("01", "FAC F001-0000002243", "FREFRA S.A.C.\n20609165066", "S/ 318.00\n(+31.80 P)", "ATP / NAC", "7", "IGV 10.5% + Recargo S/ 21.47. Propina fuera de compras."),
        ("02A", "FAC FFL1-0002517", "CABIFY LOGISTICS\n20612973335", "S/ 328.39", "GEN / XML", "8", "Recurrente (1/2). Con XML + PDF. Tasa 18%, TR."),
        ("02B", "FAC FE01-00000111", "CABIFY LOGISTICS\n20612973335", "S/ 4,744.07", "GEN / NAC", "9", "Recurrente (2/2). Con PDF. Mismo emisor para aprendizaje."),
        ("05", "FAC F004-00010092", "TOYOTA DEL PERU\n20100132592", "US$ 835.44", "MOV / NAC", "10", "Validado (VA). Tasa 18%. Aplica T/C oficial SUNAT."),
        ("06", "FAC E001-00000702", "DIDACTICA CREATIVA\n20603222921", "S/ 1,168.20", "GEN / XML", "11", "Al crédito (LC). Vencimiento informado: 24/10/2026."),
        ("07", "FAC FE01-00000112", "CABIFY LOGISTICS\n20612973335", "S/ 8,608.90", "GEN / NAC", "16", "Detracción 4% (cód. 027 S/ 344.36, cuenta BN). Crédito diferido."),
        ("08", "BV B003-00069810", "Pokeboss Pokeboss\n20603962801", "S/ 160.00", "GEN / NAC", "17", "Boleta tipo 03. Destino RCS sin crédito fiscal."),
        ("10", "FAC F041-00027050", "INVERSIONES BRADE\n20303972821", "S/ 80.80", "ATP / NAC", "12", "Proveedor nuevo no registrado. Valida contra padrón SUNAT."),
        ("11", "CND E080-10DBVA", "Microsoft Corp.\n91-1144442 (US)", "US$ 48.53", "GEN / EXT", "13", "Exterior no domiciliado (tipo 91). Inafecto. Tax ID TAX."),
        ("12", "FAC F003-00003165", "Calua S.A.C.\n20607924946", "S/ 91.60", "GEN / NAC", "14", "Multilínea: L1 Gravada (S/ 60.00) + L2 Inafecta (S/ 20.80)."),
        ("13", "FAC FFL1-0002517", "CABIFY LOGISTICS\n20612973335", "S/ 328.39", "GEN / XML", "15", "Duplicado: Mismo comprobante con UUID distinto. Alerta detectada.")
    ]

    for idx, row_data in enumerate(cases_rows):
        row = table.add_row()
        bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.width = col_widths[i]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if i in [0, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif i == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if i == 5 and val == "10":
                r.font.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 5. DETALLE TÉCNICO DE CASOS DESTACADOS
    h_detail = doc.add_paragraph()
    h_detail.paragraph_format.space_before = Pt(8)
    h_detail.paragraph_format.space_after = Pt(4)
    r_hdet = h_detail.add_run("3. Detalle Técnico de Casuísticas Relevantes")
    r_hdet.font.bold = True
    r_hdet.font.size = Pt(13)
    r_hdet.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    details = [
        ("Caso 01 — Restaurante con Recargo al Consumo y Propina:",
         "Comprobante FREFRA S.A.C. (F001-0000002243). Base gravada S/ 268.35, IGV al 10.5% (S/ 28.18) y Recargo al Consumo (S/ 21.47), totalizando S/ 318.00 en el comprobante fiscal. La propina de S/ 31.80 fue enviada en su campo independiente con voucher adjunto. El ERP registró la observación exacta: 'Propina fuera del comprobante: no entra al registro de compras', cumpliendo la Quinta Disposición Complementaria del Decreto Ley N° 25988."),

        ("Caso 05 — Factura en Dólares (USD):",
         "Factura electrónica Toyota del Perú (F004-00010092) por US$ 835.44 (Base US$ 708.00 + IGV US$ 127.44). Aceptada en estado Validado (VA). El ERP toma el tipo de cambio venta oficial SUNAT publicado a la fecha de emisión (16/09/2026) sin depender del integrador."),

        ("Caso 06 — Compra al Crédito con Vencimiento:",
         "Factura Didáctica Creativa (E001-00000702) por S/ 1,168.20. Se declaró forma_pago 'LC' (crédito) y fecha_vencimiento '2026-10-24' extraída del XML UBL, permitiendo su incorporación inmediata a la cuenta corriente del proveedor y programación de pagos."),

        ("Caso 07 — Factura Sujeta a Detracción SPOT:",
         "Factura Cabify Logistics (FE01-00000112) por S/ 8,608.90. Se informó bloque de detracción código 027 (Servicio de transporte de carga), porcentaje 4.00% (S/ 344.36), cuenta BN y valor referencial. El ERP registró la anotación tributaria: 'Detracción sin depósito acreditado; el crédito fiscal se difiere al período del depósito'."),

        ("Caso 08 — Boleta de Venta sin Crédito Fiscal:",
         "Boleta Pokeboss (B003-00069810) por S/ 160.00. El ERP la clasificó automáticamente hacia destino 'RCS' (Registro de Compras sin Crédito Fiscal), con credito_fiscal 'N' y deducible 'S'."),

        ("Caso 11 — Comprobante de No Domiciliado (EXT):",
         "Invoice Microsoft Corporation (E080-10DBVA) por US$ 48.53. Declarado en escenario EXT con tipo '91', importe clasificado íntegramente como inafecto (sin IGV) e identificado con tipo_documento TAX y país US."),

        ("Caso 13 — Control de Idempotencia y Duplicados:",
         "Se reenvió la factura Cabify FFL1-0002517 con UUID diferente ('RG-TEST-13-DUPLICADO-FFL1'). El ERP detectó la coincidencia por emisor, tipo, serie y correlativo, generando el comprobante con la observación: 'El mismo comprobante ya fue recibido antes por esta API'."
        )
    ]

    for title, text in details:
        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(4)
        p_d.paragraph_format.space_after = Pt(4)
        r_dt = p_d.add_run(f"• {title} ")
        r_dt.font.bold = True
        r_dt.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)
        r_dx = p_d.add_run(text)
        r_dx.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 6. CASOS OMITIDOS
    h_omit = doc.add_paragraph()
    h_omit.paragraph_format.space_before = Pt(8)
    h_omit.paragraph_format.space_after = Pt(4)
    r_homit = h_omit.add_run("4. Casuísticas Omitidas (Según Criterio del Alcance)")
    r_homit.font.bold = True
    r_homit.font.size = Pt(13)
    r_homit.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    omits = [
        ("Recibo de Servicios Públicos (Tipo 14):",
         "En las rendiciones de gastos y buzón de proveedores de setiembre/octubre no se registraron recibos de luz, agua o telefonía (Sedapal, Enel, Claro, etc.)."),

        ("Documento Bancario por Comisiones (Tipo 13):",
         "En el sistema se gestionan vouchers de transferencias/desembolsos (pagos a colaboradores), pero no se cuenta con estados de cuenta o liquidaciones de comisiones bancarias tipo 13 como gasto rendido."),

        ("Detracción con Constancia de Depósito Acreditada (7B):",
         "Se cubrió el Caso 7A (factura con cálculo SPOT y cuenta BN). La variante con constancia de pago acreditada se omitió al no disponerse del ticket de depósito bancario asociado."),

        ("Nota de Crédito (Tipo 07):",
         "No existieron notas de crédito reales emitidas por proveedores en el período analizado.")
    ]

    for title, text in omits:
        p_o = doc.add_paragraph()
        p_o.paragraph_format.space_before = Pt(3)
        p_o.paragraph_format.space_after = Pt(3)
        r_ot = p_o.add_run(f"– {title} ")
        r_ot.font.bold = True
        r_ot.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        r_ox = p_o.add_run(text)
        r_ox.font.size = Pt(9.5)

    # 7. CIERRE
    p_close = doc.add_paragraph()
    p_close.paragraph_format.space_before = Pt(16)
    r_cl = p_close.add_run("Todos los envíos fueron validados con los archivos adjuntos exigidos y se encuentran disponibles para la sesión de trabajo con el equipo de Contabilidad.")
    r_cl.font.italic = True
    r_cl.font.size = Pt(9.5)

    out_file = r"c:\Users\Andree\Desktop\Bliss Project\PROYECTOS SOFTWARE\Rindegastos\docs\INFORME_ENVIO_CASOS_ERP_SEAFACIL.docx"
    os.makedirs(os.path.dirname(out_file), exist_ok=True)
    doc.save(out_file)
    print(f"Documento Word generado exitosamente en: {out_file}")

if __name__ == "__main__":
    create_report()
